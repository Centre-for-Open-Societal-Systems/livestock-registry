#!/usr/bin/env python3
"""Build the livestock form's location lists from the ETH-OAN pack geography.

    python odk/tools/build_location_media.py --pack ../openg2p-data/packs/ETH-OAN
    python odk/tools/build_location_media.py --pack ... --check   # exit 1 if stale

The pack (Centre-for-Open-Societal-Systems/openg2p-data, packs/ETH-OAN) is what
New Dev Master Data is seeded from, so the choice names are the Master Data geo
unit ids (P-codes: region ET01, zone ET0101, woreda ET010101, kebele
ET01010101) and the labels are the unit names. The form's four questions are
region_code, zone_code, woreda_code and kebele_code, plus a hidden
location_source = mds, which is what odk_to_livestock_intake.j2 reads to build
the Master Data hierarchy. The lists are attached to the form in Central as
external CSVs (region_odk.csv ... kebele_odk.csv), with a parent column that the
cascading choice_filter uses.

Labels are made readable: underscores become spaces and runs of spaces collapse; places that still share a name under one parent get their code appended.
The codes are never changed. Geography is per environment: rebuild the lists
from the environment's Master Data export when it was not seeded from the pack.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from collections import Counter
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
FORM = REPO / "odk/livestock_registry.xlsx"
MEDIA = REPO / "odk/media"
VERSION = "v7"
LEVELS = {"l1": "region", "l2": "zone", "l3": "woreda", "l4": "kebele"}
PARENT = {"region": None, "zone": "region", "woreda": "zone", "kebele": "woreda"}
LISTS = {"region": "region_odk.csv", "zone": "zone_odk.csv", "woreda": "woreda_odk.csv", "kebele": "kebele_odk.csv"}


def readable(name: str) -> str:
    return re.sub(r"\s+", " ", name.replace("_", " ")).strip()


def load(pack: Path) -> dict[str, list[dict]]:
    values = json.loads((pack / "values.json").read_text(encoding="utf-8"))
    units: dict[str, list[dict]] = {level: [] for level in LEVELS.values()}
    for unit in values:
        level = LEVELS.get(unit["level_id"])
        if level:
            units[level].append(unit)
    known: dict[str, set[str]] = {level: {u["level_value_id"] for u in rows} for level, rows in units.items()}
    for level, rows in units.items():
        if len(known[level]) != len(rows):
            raise ValueError("duplicate " + level + " codes in the pack")
    for level, rows in units.items():
        parent = PARENT[level]
        for unit in rows:
            label = readable(unit.get("level_value_mnemonic") or unit["display_name"])
            if not label:
                raise ValueError("empty label for " + unit["level_value_id"])
            if parent and unit["parent_level_value_id"] not in known[parent]:
                raise ValueError("missing parent for " + unit["level_value_id"])
    return units


def labelled(units: dict[str, list[dict]], level: str) -> list[tuple[str, str, str]]:
    """(code, label, parent code) per unit. Two places under one parent that share
    a name get their code appended, so the enumerator can tell them apart."""
    rows = [(u["level_value_id"], readable(u.get("level_value_mnemonic") or u["display_name"]),
             u["parent_level_value_id"] if PARENT[level] else "") for u in units[level]]
    seen = Counter((parent, label) for _, label, parent in rows)
    return sorted(((code, f"{label} ({code})" if seen[(parent, label)] > 1 else label, parent)
                   for code, label, parent in rows))


def csv_text(units: dict[str, list[dict]], level: str) -> str:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    parent = PARENT[level]
    writer.writerow(["name", "label"] + ([parent] if parent else []))
    for code, label, parent_code in labelled(units, level):
        writer.writerow([code, label] + ([parent_code] if parent else []))
    return out.getvalue()


def update_form(workbook) -> bool:
    """region_id ... kebele_id become *_code, with a hidden location_source."""
    sheet = workbook["survey"]
    col = {c.value: c.column for c in sheet[1] if c.value}
    changed = False
    kebele_row = None  # the row of the last location question; location_source follows it
    names = {sheet.cell(r, col["name"]).value for r in range(2, sheet.max_row + 1)}
    for r in range(2, sheet.max_row + 1):
        name = sheet.cell(r, col["name"]).value
        for level in LEVELS.values():
            if name == level + "_id":
                sheet.cell(r, col["name"]).value = level + "_code"
                changed = True
                name = level + "_code"
            if name == level + "_code":
                filt = sheet.cell(r, col["choice_filter"])
                if PARENT[level] and filt.value != f"{PARENT[level]}=${{{PARENT[level]}_code}}":
                    filt.value = f"{PARENT[level]}=${{{PARENT[level]}_code}}"
                    changed = True
                if level == "kebele":
                    kebele_row = r
    if "location_source" not in names:
        if kebele_row is None:
            raise ValueError("the form has no kebele_code or kebele_id question to place location_source after")
        sheet.insert_rows(kebele_row + 1)
        sheet.cell(kebele_row + 1, col["type"], "calculate")
        sheet.cell(kebele_row + 1, col["name"], "location_source")
        sheet.cell(kebele_row + 1, col["calculation"], "'mds'")
        changed = True
    settings = workbook["settings"]
    header = {c.value: c.column for c in settings[1] if c.value}
    if settings.cell(2, header["version"]).value != VERSION:
        settings.cell(2, header["version"]).value = VERSION
        changed = True
    return changed


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pack", type=Path, required=True)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args(argv)
    units = load(args.pack)
    media = {MEDIA / LISTS[level]: csv_text(units, level) for level in LISTS}
    workbook = openpyxl.load_workbook(FORM)
    stale = [p.name for p, text in media.items() if not p.exists() or p.read_bytes().decode("utf-8") != text]
    form_changed = update_form(workbook)
    if form_changed:
        stale.append(FORM.name)
    if args.check:
        if stale:
            print("stale, rerun odk/tools/build_location_media.py: " + ", ".join(stale), file=sys.stderr)
            return 1
        print("location lists and form match the pack")
        return 0
    for path, text in media.items():
        path.write_bytes(text.encode("utf-8"))
    if form_changed:
        workbook.save(FORM)
    print("updated: " + (", ".join(stale) or "nothing"))
    print({level: len(units[level]) for level in LISTS})
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
